# Test patch: update PO Master & PO Detail
import openpyxl
from openpyxl.styles import Alignment
import os
import config

def update_po_master_detail(wb, no_po, inv_no, qty_m3, grand_total, tgl_invoice):
    # === PO Master: tambah Used M3 ===
    ws_master = wb["PO Master"]
    for rm in range(2, ws_master.max_row + 2):
        if ws_master.cell(row=rm, column=1).value == no_po:
            cur = ws_master.cell(row=rm, column=6).value or 0
            if isinstance(cur, (int, float)):
                ws_master.cell(row=rm, column=6).value = round(cur + qty_m3, 4)
            break

    # === PO Detail: tambah/update baris invoice ===
    ws_detail = wb["PO Detail"]
    inv_row = None
    next_empty = None
    for rd in range(3, ws_detail.max_row + 2):
        v = ws_detail.cell(row=rd, column=1).value
        if v is None:
            if next_empty is None:
                next_empty = rd
            break
        if str(v).strip() == str(inv_no).strip():
            inv_row = rd
            break

    if inv_row:
        # Multi-BAP: invoice sudah ada, tambah qty & total
        cur_qty = ws_detail.cell(row=inv_row, column=3).value or 0
        cur_total = ws_detail.cell(row=inv_row, column=4).value or 0
        ws_detail.cell(row=inv_row, column=3).value = round(cur_qty + qty_m3, 4)
        ws_detail.cell(row=inv_row, column=4).value = cur_total + grand_total
    else:
        # Invoice baru
        rd_new = next_empty or (ws_detail.max_row + 1)
        ws_detail.cell(row=rd_new, column=1).value = inv_no
        ws_detail.cell(row=rd_new, column=2).value = tgl_invoice
        ws_detail.cell(row=rd_new, column=3).value = qty_m3
        ws_detail.cell(row=rd_new, column=3).number_format = "#,##0.00"
        ws_detail.cell(row=rd_new, column=4).value = grand_total
        ws_detail.cell(row=rd_new, column=4).number_format = "#,##0"
        ws_detail.cell(row=rd_new, column=5).value = None
        ws_detail.cell(row=rd_new, column=6).value = f'=IF(E{rd_new}="","BELUM","LUNAS")'
        ws_detail.cell(row=rd_new, column=6).alignment = Alignment(horizontal="center")

# Test dengan data invoice 010
EXCEL = config.d("PO_Tracker_KKS.xlsx")
wb = openpyxl.load_workbook(EXCEL)

# Ambil data invoice 010 dari sheet Pengiriman
ws_p = wb["Pengiriman"]
inv_no = None
for rx in range(2, ws_p.max_row + 1):
    v = ws_p.cell(row=rx, column=5).value
    if v and "010" in str(v):
        inv_no = v
        no_po = ws_p.cell(row=rx, column=3).value
        qty_m3 = ws_p.cell(row=rx, column=7).value or 0
        tgl_inv = ws_p.cell(row=rx, column=2).value
        print(f"Found: inv={inv_no}, po={no_po}, qty={qty_m3}, tgl={tgl_inv}")
        break

if inv_no:
    # Hitung grand total dari Rp/m3
    rp_m3 = ws_p.cell(row=rx, column=8).value or 0
    sub = qty_m3 * rp_m3
    dpp = sub * 11/12
    ppn = dpp * 0.12
    grand = sub + ppn
    print(f"Grand total: {grand:,.0f}")
    update_po_master_detail(wb, no_po, inv_no, qty_m3, grand, tgl_inv)
    wb.save(EXCEL)
    print("Saved OK")
else:
    print("Invoice 010 tidak ditemukan di Pengiriman")
